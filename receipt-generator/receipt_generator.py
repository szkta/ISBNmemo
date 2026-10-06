#!/usr/bin/env python3
"""Excel の一覧から領収書 PDF を一括発行するツール。

使い方:
    # 1. 入力用の Excel テンプレートを作成
    python receipt_generator.py init receipts.xlsx

    # 2. Excel に宛名・金額・但し書きなどを記入して保存

    # 3. 領収書 PDF を発行
    python receipt_generator.py issue receipts.xlsx -o output
"""

from __future__ import annotations

import argparse
import datetime as dt
import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib.colors import HexColor, black
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

DATA_SHEET = "領収書データ"
ISSUER_SHEET = "発行者情報"

# 列見出し → 内部キー。列の並び順は問わず、見出し名で読み取る。
COLUMNS = [
    ("発行日", "date", 14),
    ("領収書番号", "number", 16),
    ("宛名", "recipient", 28),
    ("敬称", "honorific", 8),
    ("金額(税込)", "amount", 14),
    ("但し書き", "description", 30),
    ("税率(%)", "tax_rate", 9),
    ("支払方法", "payment", 12),
    ("備考", "note", 24),
    ("発行済", "issued", 10),
]
REQUIRED = ("recipient", "amount")

ISSUER_FIELDS = [
    ("発行者名", "name", "株式会社サンプル商事"),
    ("郵便番号", "postal", "〒100-0001"),
    ("住所", "address", "東京都千代田区千代田1-1-1"),
    ("電話番号", "tel", "TEL 03-0000-0000"),
    ("登録番号", "registration", "T1234567890123"),
    ("担当者", "staff", ""),
    ("印影画像", "seal_image", ""),
]

# 印紙税法上、売上代金の受取書は 5 万円以上で収入印紙が必要
STAMP_DUTY_THRESHOLD = 50_000

FONT_NAME = "JPFont"


# --------------------------------------------------------------------------
# Excel テンプレート作成
# --------------------------------------------------------------------------

def create_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (label, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"

    today = dt.date.today()
    samples = [
        (today, "", "山田 太郎", "様", 11000, "お品代として", 10, "現金", "", ""),
        (today, "", "株式会社テスト", "御中", 64800, "コンサルティング料として", 10, "銀行振込", "", ""),
        (today, "", "鈴木 花子", "様", 3240, "飲食代として", 8, "クレジットカード", "", ""),
    ]
    for r, row in enumerate(samples, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value).border = border

    col_of = {key: idx for idx, (_, key, _) in enumerate(COLUMNS, start=1)}
    letter = lambda key: ws.cell(row=1, column=col_of[key]).column_letter  # noqa: E731
    for r in range(2, 501):
        ws[f"{letter('date')}{r}"].number_format = "yyyy/mm/dd"
        ws[f"{letter('amount')}{r}"].number_format = "#,##0"

    dv_hon = DataValidation(type="list", formula1='"様,御中,殿"', allow_blank=True)
    dv_rate = DataValidation(type="list", formula1='"10,8,0"', allow_blank=True)
    dv_pay = DataValidation(
        type="list", formula1='"現金,銀行振込,クレジットカード,電子マネー,その他"', allow_blank=True
    )
    for dv, key in ((dv_hon, "honorific"), (dv_rate, "tax_rate"), (dv_pay, "payment")):
        ws.add_data_validation(dv)
        dv.add(f"{letter(key)}2:{letter(key)}500")

    info = wb.create_sheet(ISSUER_SHEET)
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 50
    for r, (label, _, default) in enumerate(ISSUER_FIELDS, start=1):
        a = info.cell(row=r, column=1, value=label)
        a.font = Font(bold=True)
        a.fill = PatternFill("solid", fgColor="DDEBF7")
        a.border = border
        info.cell(row=r, column=2, value=default).border = border
    note_row = len(ISSUER_FIELDS) + 2
    info.cell(row=note_row, column=1, value="※ 登録番号: インボイス(適格請求書発行事業者)の登録番号。不要なら空欄。")
    info.cell(row=note_row + 1, column=1, value="※ 印影画像: 印影 PNG のパス(Excel ファイルからの相対パス可)。不要なら空欄。")

    help_ws = wb.create_sheet("使い方")
    help_ws.column_dimensions["A"].width = 100
    lines = [
        "【使い方】",
        "1. 「発行者情報」シートに自社の名称・住所などを入力します。",
        "2. 「領収書データ」シートに 1 行 = 1 枚 で領収書の内容を入力します(サンプル行は削除してください)。",
        "   ・必須: 宛名、金額(税込)",
        "   ・発行日が空欄なら発行した日付、領収書番号が空欄なら自動採番されます。",
        "   ・税率(%) は 10 / 8 / 0 (空欄なら 10)。税抜金額と消費税額は自動計算されます。",
        "   ・金額が 5 万円以上の場合は収入印紙の貼付欄が表示されます。",
        "3. 保存して閉じ、コマンドを実行します:  python receipt_generator.py issue この.xlsx",
        "4. 発行に成功した行は「発行済」列に日時が記入され、次回以降はスキップされます。",
        "   再発行したい場合は「発行済」を空欄にするか、--all オプションを付けて実行してください。",
    ]
    for r, text in enumerate(lines, start=1):
        help_ws.cell(row=r, column=1, value=text)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------
# Excel 読み込み
# --------------------------------------------------------------------------

@dataclass
class Issuer:
    name: str = ""
    postal: str = ""
    address: str = ""
    tel: str = ""
    registration: str = ""
    staff: str = ""
    seal_image: Path | None = None


@dataclass
class Receipt:
    row: int
    date: dt.date
    number: str
    recipient: str
    honorific: str
    amount: int
    description: str
    tax_rate: int
    payment: str
    note: str

    @property
    def tax(self) -> int:
        # 税込金額から内税の消費税額を算出 (1 円未満切り捨て)
        if self.tax_rate <= 0:
            return 0
        return math.floor(self.amount * self.tax_rate / (100 + self.tax_rate))

    @property
    def amount_excl_tax(self) -> int:
        return self.amount - self.tax


class InputError(Exception):
    pass


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _parse_date(value, row: int) -> dt.date | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _text(value)
    m = re.fullmatch(r"(\d{4})\D(\d{1,2})\D(\d{1,2})\D?", text)
    if not m:
        raise InputError(f"{row} 行目: 発行日 '{text}' を日付として読めません (例: 2026/10/06)")
    return dt.date(*map(int, m.groups()))


def _parse_int(value, label: str, row: int) -> int | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise InputError(f"{row} 行目: {label} は整数で入力してください ({value})")
        return int(value)
    text = re.sub(r"[,，¥￥円\s%％]", "", _text(value))
    text = text.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if not re.fullmatch(r"-?\d+", text):
        raise InputError(f"{row} 行目: {label} '{value}' を数値として読めません")
    return int(text)


def read_issuer(wb, base_dir: Path) -> Issuer:
    issuer = Issuer()
    if ISSUER_SHEET not in wb.sheetnames:
        return issuer
    ws = wb[ISSUER_SHEET]
    label_to_key = {label: key for label, key, _ in ISSUER_FIELDS}
    for label, value in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        key = label_to_key.get(_text(label))
        if not key:
            continue
        text = _text(value)
        if key == "seal_image":
            if text:
                p = Path(text).expanduser()
                issuer.seal_image = p if p.is_absolute() else base_dir / p
        else:
            setattr(issuer, key, text)
    return issuer


def read_receipts(ws, include_issued: bool) -> tuple[list[Receipt], dict[str, int]]:
    header = [_text(c.value) for c in ws[1]]
    label_to_key = {label: key for label, key, _ in COLUMNS}
    col_index: dict[str, int] = {}
    for idx, label in enumerate(header):
        key = label_to_key.get(label)
        if key:
            col_index[key] = idx
    missing = [label for label, key, _ in COLUMNS if key in REQUIRED and key not in col_index]
    if missing:
        raise InputError(f"「{DATA_SHEET}」シートに列 {', '.join(missing)} がありません")

    receipts: list[Receipt] = []
    errors: list[str] = []
    for row_cells in ws.iter_rows(min_row=2):
        row = row_cells[0].row
        get = lambda key: row_cells[col_index[key]].value if key in col_index and col_index[key] < len(row_cells) else None  # noqa: E731
        if all(_text(get(k)) == "" for k in ("recipient", "amount", "description")):
            continue  # 空行
        if not include_issued and _text(get("issued")):
            continue
        try:
            recipient = _text(get("recipient"))
            if not recipient:
                raise InputError(f"{row} 行目: 宛名が空欄です")
            amount = _parse_int(get("amount"), "金額", row)
            if amount is None or amount <= 0:
                raise InputError(f"{row} 行目: 金額は 1 以上の数値で入力してください")
            tax_rate = _parse_int(get("tax_rate"), "税率", row)
            receipts.append(
                Receipt(
                    row=row,
                    date=_parse_date(get("date"), row) or dt.date.today(),
                    number=_text(get("number")),
                    recipient=recipient,
                    honorific=_text(get("honorific")) or "様",
                    amount=amount,
                    description=_text(get("description")) or "お品代として",
                    tax_rate=10 if tax_rate is None else tax_rate,
                    payment=_text(get("payment")),
                    note=_text(get("note")),
                )
            )
        except InputError as e:
            errors.append(str(e))
    if errors:
        raise InputError("入力内容にエラーがあります:\n  " + "\n  ".join(errors))
    return receipts, col_index


def assign_numbers(receipts: list[Receipt], ws, col_index: dict[str, int]) -> None:
    """領収書番号が空欄の行に R-YYYYMMDD-NNN 形式で採番する。"""
    used = set()
    if "number" in col_index:
        for (value,) in ws.iter_rows(
            min_row=2, min_col=col_index["number"] + 1, max_col=col_index["number"] + 1, values_only=True
        ):
            if value:
                used.add(_text(value))
    for r in receipts:
        if r.number:
            continue
        prefix = f"R-{r.date:%Y%m%d}-"
        seq = 1
        while f"{prefix}{seq:03d}" in used:
            seq += 1
        r.number = f"{prefix}{seq:03d}"
        used.add(r.number)


# --------------------------------------------------------------------------
# PDF 描画
# --------------------------------------------------------------------------

# PDF に埋め込む日本語フォントの候補 (TrueType アウトラインのもの)
FONT_CANDIDATES = [
    Path(__file__).resolve().parent / "fonts" / "ipaexg.ttf",
    Path("C:/Windows/Fonts/msgothic.ttc"),
    Path("C:/Windows/Fonts/meiryo.ttc"),
    Path("C:/Windows/Fonts/YuGothM.ttc"),
    Path("/Library/Fonts/Arial Unicode.ttf"),
    Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    Path("/usr/share/fonts/opentype/ipaexfont-gothic/ipaexg.ttf"),
    Path("/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf"),
    Path("/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"),
]


def _try_register_ttf(path: Path) -> bool:
    try:
        pdfmetrics.registerFont(TTFont("JPFont", str(path), subfontIndex=0))
        return True
    except Exception:  # noqa: BLE001 - CFF 形式など読めないフォントは次の候補へ
        return False


def register_font(font_path: str | None) -> None:
    global FONT_NAME
    FONT_NAME = "JPFont"
    if font_path:
        if not _try_register_ttf(Path(font_path)):
            raise InputError(f"フォント {font_path} を読み込めません (TrueType 形式の .ttf/.ttc を指定してください)")
        return
    for candidate in FONT_CANDIDATES:
        if candidate.exists() and _try_register_ttf(candidate):
            return
    # 埋め込み可能なフォントが無い場合は PDF ビューア側の日本語フォントに頼る CID フォントを使う
    print(
        "警告: 埋め込み用の日本語フォントが見つからないため CID フォントを使用します。"
        " 環境によっては文字が表示されません (--font で .ttf を指定してください)。",
        file=sys.stderr,
    )
    FONT_NAME = "HeiseiKakuGo-W5"
    pdfmetrics.registerFont(UnicodeCIDFont(FONT_NAME))


def _fmt_date(d: dt.date) -> str:
    return f"{d.year}年{d.month}月{d.day}日"


def _fit_font_size(text: str, max_width: float, size: float, min_size: float = 7) -> float:
    while size > min_size and pdfmetrics.stringWidth(text, FONT_NAME, size) > max_width:
        size -= 0.5
    return size


def draw_receipt(c: canvas.Canvas, r: Receipt, issuer: Issuer, top: float, title: str, is_copy: bool = False) -> None:
    """A4 の上端 top から下に、幅いっぱい・高さ約 140mm の領収書 1 枚を描画する。"""
    page_w, _ = A4
    left, right = 18 * mm, page_w - 18 * mm
    width = right - left
    accent = HexColor("#1F4E78")

    # タイトル
    c.setFillColor(accent)
    c.setFont(FONT_NAME, 22)
    c.drawCentredString(page_w / 2, top - 16 * mm, title)
    c.setFillColor(black)
    title_w = pdfmetrics.stringWidth(title, FONT_NAME, 22)
    c.setStrokeColor(accent)
    c.setLineWidth(1)
    c.line(page_w / 2 - title_w / 2 - 6 * mm, top - 18.5 * mm, page_w / 2 + title_w / 2 + 6 * mm, top - 18.5 * mm)

    # 番号・日付
    c.setFont(FONT_NAME, 9.5)
    c.drawRightString(right, top - 8 * mm, f"No. {r.number}")
    c.drawRightString(right, top - 13 * mm, f"発行日 {_fmt_date(r.date)}")

    # 宛名
    y = top - 33 * mm
    name_w = width * 0.62
    suffix = f" {r.honorific}"
    size = _fit_font_size(r.recipient + suffix, name_w, 16)
    c.setFont(FONT_NAME, size)
    c.drawString(left + 2 * mm, y, r.recipient + suffix)
    c.setStrokeColor(black)
    c.setLineWidth(0.8)
    c.line(left, y - 2.5 * mm, left + name_w, y - 2.5 * mm)

    # 金額ボックス
    box_top = y - 9 * mm
    box_h = 18 * mm
    box_l, box_w = left + 15 * mm, width - 30 * mm
    c.setFillColor(HexColor("#EEF3F9"))
    c.setStrokeColor(accent)
    c.setLineWidth(1.2)
    c.rect(box_l, box_top - box_h, box_w, box_h, stroke=1, fill=1)
    c.setFillColor(black)
    c.setFont(FONT_NAME, 11)
    c.drawString(box_l + 6 * mm, box_top - box_h / 2 - 1.5 * mm, "金額")
    c.setFont(FONT_NAME, 24)
    c.drawCentredString(box_l + box_w / 2 + 6 * mm, box_top - box_h / 2 - 3 * mm, f"¥{r.amount:,}-")
    c.setFont(FONT_NAME, 9)
    c.drawRightString(box_l + box_w - 4 * mm, box_top - box_h / 2 - 1.5 * mm, "(税込)")

    # 但し書き
    y = box_top - box_h - 9 * mm
    c.setFont(FONT_NAME, 11)
    desc = f"但し　{r.description}"
    c.setFont(FONT_NAME, _fit_font_size(desc, width * 0.75, 11))
    c.drawString(left + 2 * mm, y, desc)
    c.setFont(FONT_NAME, 11)
    c.drawString(left + 2 * mm, y - 6.5 * mm, "上記正に領収いたしました。")

    # 内訳 (左下)
    y_tbl = y - 16 * mm
    rows = [("税抜金額", f"¥{r.amount_excl_tax:,}"), (f"消費税額 ({r.tax_rate}%)", f"¥{r.tax:,}")]
    if r.tax_rate == 8:
        rows[1] = ("消費税額 (8% 軽減税率)", f"¥{r.tax:,}")
    if r.payment:
        rows.append(("お支払方法", r.payment))
    tbl_w, row_h = 80 * mm, 6.5 * mm
    c.setLineWidth(0.5)
    c.setStrokeColor(HexColor("#7F7F7F"))
    c.setFont(FONT_NAME, 9)
    c.drawString(left, y_tbl + 1.5 * mm, "内訳")
    for i, (k, v) in enumerate(rows):
        ry = y_tbl - (i + 1) * row_h
        c.rect(left, ry, tbl_w, row_h)
        c.line(left + tbl_w * 0.55, ry, left + tbl_w * 0.55, ry + row_h)
        c.drawString(left + 2 * mm, ry + 2 * mm, k)
        c.drawRightString(left + tbl_w - 2 * mm, ry + 2 * mm, v)
    if r.note:
        ry = y_tbl - (len(rows) + 1) * row_h + 1 * mm
        note = f"備考: {r.note}"
        c.setFont(FONT_NAME, _fit_font_size(note, tbl_w, 8.5))
        c.drawString(left, ry, note)

    # 収入印紙欄 (5 万円以上。控えには不要)
    if r.amount >= STAMP_DUTY_THRESHOLD and not is_copy:
        sw = 22 * mm
        sx, sy = left + tbl_w + 8 * mm, y_tbl - sw
        c.setDash(2, 2)
        c.rect(sx, sy, sw, sw)
        c.setDash()
        c.setFont(FONT_NAME, 7.5)
        c.drawCentredString(sx + sw / 2, sy + sw / 2 + 1 * mm, "収入印紙")
        c.drawCentredString(sx + sw / 2, sy + sw / 2 - 3 * mm, "貼付欄")

    # 発行者 (右下)
    ix = right - 62 * mm
    iy = y_tbl
    c.setFont(FONT_NAME, 11.5)
    c.drawString(ix, iy, issuer.name)
    c.setFont(FONT_NAME, 8.5)
    lines = [issuer.postal, issuer.address, issuer.tel]
    if issuer.registration:
        lines.append(f"登録番号 {issuer.registration}")
    if issuer.staff:
        lines.append(f"担当 {issuer.staff}")
    ly = iy - 5.5 * mm
    for line in filter(None, lines):
        c.drawString(ix, ly, line)
        ly -= 4.5 * mm

    # 印影
    seal = 20 * mm
    seal_x, seal_y = right - seal, iy - 9 * mm
    if issuer.seal_image and issuer.seal_image.exists():
        c.drawImage(ImageReader(str(issuer.seal_image)), seal_x, seal_y, seal, seal, mask="auto",
                    preserveAspectRatio=True)


def render_page(c: canvas.Canvas, r: Receipt, issuer: Issuer, with_copy: bool) -> None:
    page_w, page_h = A4
    draw_receipt(c, r, issuer, page_h - 10 * mm, "領　収　書")
    if with_copy:
        mid = page_h / 2
        c.setDash(4, 3)
        c.setStrokeColor(HexColor("#A6A6A6"))
        c.setLineWidth(0.5)
        c.line(8 * mm, mid, page_w - 8 * mm, mid)
        c.setDash()
        c.setFont(FONT_NAME, 7)
        c.setFillColor(HexColor("#7F7F7F"))
        c.drawCentredString(page_w / 2, mid + 1.5 * mm, "切り取り線")
        c.setFillColor(black)
        draw_receipt(c, r, issuer, mid - 10 * mm, "領収書（控）", is_copy=True)
    c.showPage()


def _safe_filename(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text).strip("_")[:40]


def issue(
    excel_path: Path,
    out_dir: Path,
    include_issued: bool,
    with_copy: bool,
    combined: bool,
    mark_issued: bool,
    font_path: str | None,
) -> list[Path]:
    wb = load_workbook(excel_path)
    if DATA_SHEET not in wb.sheetnames:
        raise InputError(f"シート「{DATA_SHEET}」が見つかりません")
    ws = wb[DATA_SHEET]
    issuer = read_issuer(wb, excel_path.parent)
    if issuer.seal_image and not issuer.seal_image.exists():
        print(f"警告: 印影画像が見つかりません: {issuer.seal_image}", file=sys.stderr)

    receipts, col_index = read_receipts(ws, include_issued)
    if not receipts:
        print("発行対象の行がありません (発行済の行を再発行するには --all を指定してください)。")
        return []
    assign_numbers(receipts, ws, col_index)

    register_font(font_path)
    out_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for r in receipts:
        path = out_dir / f"{_safe_filename(r.number)}_{_safe_filename(r.recipient)}.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.setTitle(f"領収書 {r.number}")
        c.setAuthor(issuer.name)
        render_page(c, r, issuer, with_copy)
        c.save()
        outputs.append(path)

    if combined and len(receipts) > 1:
        path = out_dir / f"領収書_一括_{dt.datetime.now():%Y%m%d_%H%M%S}.pdf"
        c = canvas.Canvas(str(path), pagesize=A4)
        c.setTitle("領収書 一括")
        for r in receipts:
            render_page(c, r, issuer, with_copy)
        c.save()
        outputs.append(path)

    if mark_issued:
        _write_back(wb, ws, receipts, col_index)
        try:
            wb.save(excel_path)
        except PermissionError:
            print(
                f"警告: {excel_path} に書き込めませんでした (Excel で開いたままの可能性があります)。"
                " 領収書番号・発行済の記録は保存されていません。",
                file=sys.stderr,
            )
    return outputs


def _write_back(wb, ws, receipts: list[Receipt], col_index: dict[str, int]) -> None:
    def column_for(key: str) -> int:
        if key in col_index:
            return col_index[key] + 1
        label = next(lbl for lbl, k, _ in COLUMNS if k == key)
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=label)
        col_index[key] = col - 1
        return col

    num_col = column_for("number")
    issued_col = column_for("issued")
    date_col = column_for("date")
    stamp = dt.datetime.now().strftime("%Y/%m/%d %H:%M")
    for r in receipts:
        ws.cell(row=r.row, column=num_col, value=r.number)
        ws.cell(row=r.row, column=issued_col, value=stamp)
        date_cell = ws.cell(row=r.row, column=date_col)
        if date_cell.value in (None, ""):
            date_cell.value = r.date
            date_cell.number_format = "yyyy/mm/dd"


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Excel から領収書 PDF を自動発行します。")
    sub = parser.add_subparsers(dest="command", required=True)

    p_init = sub.add_parser("init", help="入力用の Excel テンプレートを作成する")
    p_init.add_argument("excel", type=Path, help="作成する .xlsx のパス")
    p_init.add_argument("--force", action="store_true", help="既存ファイルを上書きする")

    p_issue = sub.add_parser("issue", help="Excel の内容から領収書 PDF を発行する")
    p_issue.add_argument("excel", type=Path, help="入力 .xlsx のパス")
    p_issue.add_argument("-o", "--out", type=Path, default=Path("output"), help="PDF の出力先 (既定: ./output)")
    p_issue.add_argument("--all", action="store_true", help="発行済の行も含めて再発行する")
    p_issue.add_argument("--no-copy", action="store_true", help="控えを付けない (1 ページに領収書のみ)")
    p_issue.add_argument("--no-combined", action="store_true", help="全件をまとめた PDF を作らない")
    p_issue.add_argument("--no-mark", action="store_true", help="Excel に領収書番号・発行済を書き戻さない")
    p_issue.add_argument("--font", help="使用する日本語 TrueType フォント (.ttf) のパス。PDF へ埋め込まれます")

    args = parser.parse_args(argv)
    try:
        if args.command == "init":
            if args.excel.exists() and not args.force:
                print(f"{args.excel} は既に存在します (上書きする場合は --force)。", file=sys.stderr)
                return 1
            create_template(args.excel)
            print(f"テンプレートを作成しました: {args.excel}")
            return 0

        if not args.excel.exists():
            print(f"ファイルが見つかりません: {args.excel}", file=sys.stderr)
            return 1
        outputs = issue(
            args.excel,
            args.out,
            include_issued=args.all,
            with_copy=not args.no_copy,
            combined=not args.no_combined,
            mark_issued=not args.no_mark,
            font_path=args.font,
        )
        for p in outputs:
            print(f"発行: {p}")
        if outputs:
            print(f"合計 {len([p for p in outputs if '一括' not in p.name])} 件の領収書を発行しました。")
        return 0
    except InputError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
